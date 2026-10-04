# -*- coding: utf-8 -*-
"""
model.py – cała logika Skarbnika (BEZ okien), żeby dało się ją testować osobno.

Dlaczego osobny plik:
- Przeliczenia (sumy, salda, statusy, ściągalność) są zwykłym kodem Pythona, a nie
  formułami arkusza. Dzięki temu aplikacja działa bez MS Office i nie ma problemu
  z przeliczaniem formuł – wynik jest zawsze świeży po każdej zmianie.
- Dane trzymamy w jednym pliku JSON (czytelny, łatwy do skopiowania jako backup).
"""
import copy
import datetime
import json
import os
import shutil

DOMYSLNE_KATEGORIE = ["Wycieczki", "Prezenty i drobiazgi", "Materiały naukowe",
                      "Imprezy klasowe", "Fundusz klasowy", "Inne"]
DOMYSLNE_FORMY = ["Gotówka", "Przelew"]
DOMYSLNE_SKLADKI = [("Składka roczna", 150.0), ("Wycieczka jesienna", 60.0),
                    ("Dzień Edukacji Narodowej", 15.0), ("Mikołajki", 20.0),
                    ("Wycieczka wiosenna", 80.0)]


# ---------------------------------------------------------------- dane
def nowe():
    """Pusty zestaw danych z domyślnymi słownikami (pierwsze uruchomienie)."""
    return {
        "wersja": 1,
        "parametry": {"klasa": "Klasa", "rok": "", "skarbnik": "",
                      "saldo_poczatkowe": 0.0, "inne_przychody": 0.0},
        "kategorie_wydatkow": list(DOMYSLNE_KATEGORIE),
        "formy_platnosci": list(DOMYSLNE_FORMY),
        "skladki": [{"nazwa": n, "stawka": s} for n, s in DOMYSLNE_SKLADKI],
        "uczniowie": [],
        "wydatki": [],
    }


def przyklad():
    """Dane przykładowe – żeby od razu zobaczyć, jak program działa."""
    d = nowe()
    d["parametry"] = {"klasa": "Klasa 5a", "rok": "2026/2027", "skarbnik": "Ewa Mazur",
                      "saldo_poczatkowe": 200.0, "inne_przychody": 100.0}
    lista = [("Anna", "Kowalska", [150, 60, 15, 20, 80]), ("Jan", "Nowak", [150, 60, 15, 20, 40]),
             ("Piotr", "Wiśniewski", [150, 60, 15, 20, 0]), ("Maria", "Wójcik", [150, 60, 15, 20, 80]),
             ("Tomasz", "Kamiński", [150, 0, 15, 0, 0]), ("Zofia", "Lewandowska", [150, 60, 15, 20, 100]),
             ("Michał", "Zieliński", [100, 60, 15, 20, 0]), ("Julia", "Szymańska", [150, 60, 15, 20, 80])]
    for i, (im, na, w) in enumerate(lista, 1):
        d["uczniowie"].append({"nr": i, "imie": im, "nazwisko": na, "znizka": 0.0,
                               "wplaty": [float(x) for x in w]})
    d["wydatki"] = [
        {"data": "2026-09-12", "dowod": "PAR/001/2026", "opis": "Bilety wstępu do muzeum",
         "kategoria": "Wycieczki", "kwota": 540.0, "forma": "Przelew", "osoba": "Ewa Mazur", "uwagi": ""},
        {"data": "2026-09-20", "dowod": "Par. 4521", "opis": "Zeszyty i pomoce naukowe",
         "kategoria": "Materiały naukowe", "kwota": 185.5, "forma": "Gotówka", "osoba": "Ewa Mazur", "uwagi": ""},
    ]
    return d


def normalizuj(d):
    """Dopasowuje listy wpłat uczniów do liczby składek (po zmianie słownika).

    Wpłaty trzymamy POZYCYJNIE (i-ta wpłata = i-ta składka) – dzięki temu zmiana
    nazwy składki nie gubi wpłat. Brakujące miejsca dopełniamy zerami.
    """
    n = len(d["skladki"])
    for u in d["uczniowie"]:
        w = [float(x) for x in u.get("wplaty", [])][:n]
        u["wplaty"] = w + [0.0] * (n - len(w))
        u.setdefault("znizka", 0.0)
    return d


# ---------------------------------------------------------------- zapis / odczyt
def zapisz(d, sciezka):
    """Zapis atomowy: najpierw plik tymczasowy, potem podmiana – awaria w trakcie
    zapisu nie zniszczy poprzedniej wersji. Poprzednią wersję zostawiamy jako .bak."""
    folder = os.path.dirname(os.path.abspath(sciezka))
    os.makedirs(folder, exist_ok=True)
    tmp = sciezka + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=2)
    if os.path.exists(sciezka):
        shutil.copy2(sciezka, sciezka + ".bak")
    os.replace(tmp, sciezka)


def wczytaj(sciezka):
    with open(sciezka, "r", encoding="utf-8") as f:
        d = json.load(f)
    base = nowe()
    for k, v in base.items():          # uzupełnij brakujące klucze (kompatybilność wstecz)
        d.setdefault(k, v)
    return normalizuj(d)


# ---------------------------------------------------------------- parsowanie i formatowanie
def parse_kwota(tekst):
    """'1 234,50 zł' -> 1234.5. Puste pole = 0. Zły tekst -> ValueError."""
    if tekst is None:
        return 0.0
    t = str(tekst).replace("zł", "").replace("\u00a0", "").replace(" ", "").replace(",", ".").strip()
    if t in ("", "-"):
        return 0.0
    return round(float(t), 2)


def parse_data(tekst):
    """Przyjmuje 12.09.2026, 12-09-2026, 12/09/2026 lub 2026-09-12; zwraca ISO."""
    t = str(tekst).strip()
    for fmt in ("%d.%m.%Y", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return datetime.datetime.strptime(t, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError("Wpisz datę w formacie DD.MM.RRRR, np. 12.09.2026")


def fmt_data(iso):
    try:
        return datetime.date.fromisoformat(iso).strftime("%d.%m.%Y")
    except Exception:
        return iso or ""


def fmt_kw(v):
    """1234.5 -> '1 234,50' (polski zapis)."""
    return f"{v:,.2f}".replace(",", " ").replace(".", ",")


def fmt_zl(v):
    return fmt_kw(v) + " zł"


def fmt_wplata(v):
    return "-" if not v else fmt_kw(v)


# ---------------------------------------------------------------- obliczenia
def wymagana(u, d):
    return round(sum(s["stawka"] for s in d["skladki"]) - (u.get("znizka") or 0), 2)


def wplacona(u):
    return round(sum(u["wplaty"]), 2)


def saldo(u, d):
    # round(...,2) – chroni przed błędami zmiennoprzecinkowymi typu 0.0000001
    return round(wplacona(u) - wymagana(u, d), 2)


def status(u, d):
    s, w = saldo(u, d), wymagana(u, d)
    if s < 0:
        return "Zaległość"
    if s == 0 and w >= 0:
        return "Zapłacone"
    return "Nadpłata"


def sumy_uczniow(d):
    """Wiersze podsumowujące listy uczniów: suma, liczba wpłacających, średnia."""
    n = len(d["skladki"])
    suma, licz = [0.0] * n, [0] * n
    for u in d["uczniowie"]:
        for i, x in enumerate(u["wplaty"]):
            suma[i] += x
            licz[i] += 1 if x > 0 else 0
    wpl = [wplacona(u) for u in d["uczniowie"]]
    sr = [(suma[i] / licz[i]) if licz[i] else 0.0 for i in range(n)]
    return {"suma": [round(x, 2) for x in suma], "liczba": licz, "srednia": sr,
            "wpl_suma": round(sum(wpl), 2), "wpl_liczba": sum(1 for x in wpl if x > 0),
            "wpl_srednia": (sum(wpl) / sum(1 for x in wpl if x > 0)) if any(x > 0 for x in wpl) else 0.0,
            "wym_suma": round(sum(wymagana(u, d) for u in d["uczniowie"]), 2),
            "saldo_suma": round(sum(saldo(u, d) for u in d["uczniowie"]), 2),
            "znizki_suma": round(sum(u.get("znizka") or 0 for u in d["uczniowie"]), 2)}


def podsumowanie(d):
    p = d["parametry"]
    budzet = round(sum(wymagana(u, d) for u in d["uczniowie"]), 2)
    wplaty = round(sum(wplacona(u) for u in d["uczniowie"]), 2)
    # Zaległości = suma sald UJEMNYCH; nadpłaty nie pomniejszają długu innych uczniów
    zaleglosci = round(-sum(s for s in (saldo(u, d) for u in d["uczniowie"]) if s < 0), 2)
    wydatki = round(sum(w["kwota"] for w in d["wydatki"]), 2)
    razem = round(p["saldo_poczatkowe"] + wplaty + p["inne_przychody"], 2)
    stat = [status(u, d) for u in d["uczniowie"]]
    kat = []
    for k in d["kategorie_wydatkow"]:
        wg_formy = {f: round(sum(w["kwota"] for w in d["wydatki"]
                                 if w["kategoria"] == k and w["forma"] == f), 2) for f in d["formy_platnosci"]}
        r = round(sum(w["kwota"] for w in d["wydatki"] if w["kategoria"] == k), 2)
        kat.append({"kategoria": k, "formy": wg_formy, "razem": r,
                    "proc": (r / wydatki) if wydatki else 0.0})
    return {
        "budzet": budzet, "wplaty": wplaty, "wydatki": wydatki,
        "saldo_poczatkowe": p["saldo_poczatkowe"], "inne": p["inne_przychody"],
        "razem_srodki": razem, "saldo": round(razem - wydatki, 2),
        "zaleglosci": zaleglosci,
        "sciagalnosc": (1 - zaleglosci / budzet) if budzet else 0.0,
        "liczba": len(d["uczniowie"]), "zaplacone": stat.count("Zapłacone"),
        "zaleglosc": stat.count("Zaległość"), "nadplata": stat.count("Nadpłata"),
        "kategorie": kat,
        # kontrola: wydatki, których kategorii nie ma w słowniku
        "nieprzypisane": round(wydatki - sum(k["razem"] for k in kat), 2),
    }


def lista_zaleglosci(d):
    """Tekst do wysłania rodzicom: kto ile jeszcze winien."""
    wiersze = [f'{u["imie"]} {u["nazwisko"]} – {fmt_zl(-saldo(u, d))}'
               for u in d["uczniowie"] if saldo(u, d) < 0]
    return "\n".join(wiersze)


# ---------------------------------------------------------------- import / eksport Excela
def import_xlsx(sciezka):
    """Wczytuje dane z arkusza 'Skarbnik_Klasowy.xlsx' (kolumny wejściowe, nie formuły)."""
    from openpyxl import load_workbook
    wb = load_workbook(sciezka, data_only=True)
    s, u, r = wb["Słowniki i Parametry"], wb["Uczniowie i Wpłaty"], wb["Rejestr Wydatków"]
    d = nowe()
    zn = lambda ws, rng: [c.value for row in ws[rng] for c in row if c.value not in (None, "")]
    d["kategorie_wydatkow"] = zn(s, "A5:A12") or d["kategorie_wydatkow"]
    d["formy_platnosci"] = zn(s, "E5:E8") or d["formy_platnosci"]
    d["skladki"] = [{"nazwa": str(s[f"G{i}"].value), "stawka": float(s[f"H{i}"].value or 0)}
                    for i in range(5, 11) if s[f"G{i}"].value not in (None, "")]
    d["parametry"] = {"klasa": str(s["K5"].value or ""), "rok": str(s["K6"].value or ""),
                      "skarbnik": str(s["K7"].value or ""),
                      "saldo_poczatkowe": float(s["K8"].value or 0),
                      "inne_przychody": float(s["K9"].value or 0)}
    n = len(d["skladki"])
    for i in range(6, 36):
        if u[f"C{i}"].value in (None, ""):
            continue
        d["uczniowie"].append({
            "nr": int(u[f"A{i}"].value or len(d["uczniowie"]) + 1),
            "imie": str(u[f"B{i}"].value or ""), "nazwisko": str(u[f"C{i}"].value),
            "znizka": float(u[f"J{i}"].value or 0),
            "wplaty": [float(u.cell(i, 4 + k).value or 0) for k in range(n)]})
    for i in range(5, 305):
        if r[f"E{i}"].value in (None, ""):
            continue
        dt = r[f"A{i}"].value
        d["wydatki"].append({
            "data": dt.date().isoformat() if hasattr(dt, "date") else (parse_data(dt) if dt else ""),
            "dowod": str(r[f"B{i}"].value or ""), "opis": str(r[f"C{i}"].value or ""),
            "kategoria": str(r[f"D{i}"].value or ""), "kwota": float(r[f"E{i}"].value),
            "forma": str(r[f"F{i}"].value or ""), "osoba": str(r[f"G{i}"].value or ""),
            "uwagi": str(r[f"H{i}"].value or "")})
    return normalizuj(d)


def eksport_xlsx(d, sciezka):
    """Raport do Excela/LibreOffice (same wartości) – do wydruku lub przesłania dyrekcji."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    hdr_f = Font(name="Arial", bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", start_color="1F3864")
    base = Font(name="Arial", size=10)

    def tabela(ws, start_row, naglowki, wiersze, fmt_kol=None):
        for j, h in enumerate(naglowki, 1):
            c = ws.cell(start_row, j, h); c.font = hdr_f; c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, w in enumerate(wiersze, start_row + 1):
            for j, v in enumerate(w, 1):
                c = ws.cell(i, j, v); c.font = base
                if fmt_kol and j in fmt_kol:
                    c.number_format = fmt_kol[j]
        return start_row + len(wiersze) + 1

    money = '#,##0.00 "zł"'
    p = podsumowanie(d)
    ws = wb.active; ws.title = "Dashboard"
    ws["A1"] = f'Skarbnik klasowy – {d["parametry"]["klasa"]}  {d["parametry"]["rok"]}'
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws["A2"] = f'Stan na dzień {datetime.date.today().strftime("%d.%m.%Y")}'
    wiersze = [("Saldo początkowe", p["saldo_poczatkowe"]), ("Wpłaty składek", p["wplaty"]),
               ("Inne przychody", p["inne"]), ("Razem środki", p["razem_srodki"]),
               ("Łączne wydatki", p["wydatki"]), ("Aktualne saldo", p["saldo"]),
               ("Planowany budżet", p["budzet"]), ("Zaległości", p["zaleglosci"]),
               ("Ściągalność składek", p["sciagalnosc"])]
    kon = tabela(ws, 4, ["Pozycja", "Kwota"], wiersze, {2: money})
    ws.cell(kon - 1, 2).number_format = "0.0%"
    formy = d["formy_platnosci"]
    tabela(ws, kon + 2, ["Kategoria"] + formy + ["Razem", "% udziału"],
           [[k["kategoria"]] + [k["formy"][f] for f in formy] + [k["razem"], k["proc"]] for k in p["kategorie"]],
           {**{j: money for j in range(2, 3 + len(formy))}, 3 + len(formy): "0.0%"})
    ws.column_dimensions["A"].width = 34
    for c in "BCDE":
        ws.column_dimensions[c].width = 18

    ws = wb.create_sheet("Uczniowie i Wpłaty")
    sk = d["skladki"]
    tabela(ws, 1, ["Nr", "Imię", "Nazwisko"] + [s["nazwa"] for s in sk] +
           ["Zniżka", "Suma wymagana", "Suma wpłacona", "Saldo / Zaległość", "Status"],
           [[u["nr"], u["imie"], u["nazwisko"]] + u["wplaty"] +
            [u.get("znizka") or 0, wymagana(u, d), wplacona(u), saldo(u, d), status(u, d)]
            for u in d["uczniowie"]], {j: money for j in range(4, 9 + len(sk))})
    for col in range(1, 10 + len(sk)):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 16

    ws = wb.create_sheet("Rejestr Wydatków")
    tabela(ws, 1, ["Data", "Nr dowodu", "Opis", "Kategoria", "Kwota", "Forma płatności", "Osoba", "Uwagi"],
           [[fmt_data(w["data"]), w["dowod"], w["opis"], w["kategoria"], w["kwota"], w["forma"],
             w["osoba"], w["uwagi"]] for w in sorted(d["wydatki"], key=lambda x: x["data"])], {5: money})
    for col, wd in zip("ABCDEFGH", [12, 18, 40, 22, 14, 16, 20, 24]):
        ws.column_dimensions[col].width = wd
    wb.save(sciezka)
